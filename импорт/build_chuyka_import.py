# -*- coding: utf-8 -*-
"""
Готує Horoshop-імпорт (66 колонок) для розділу "Детектори дронів" — лінійка
товарів «Чуйка» від Blue Bird Tech.

Джерела (обидва в цій теці):
  * marketplace_feed_29_09 (1).xml  — YML-фід: назви UA/RU, описи, фото, параметри
  * ПРАЙС_ДЛЯ_ПАРТНЕРІВ_ОНОВЛЕНИЙ_0110.xlsx — АКТУАЛЬНІ ціни (мають пріоритет)

Беремо рівно 7 позицій «Чуйка» з прайсу, підтягуємо до них опис/фото/параметри
з фіду. Ціна — з прайсу (свіжіша). Без варіацій — кожен товар окрема картка.

Фото: blue-bird.tech віддає 403 без Referer (hotlink-захист), тому Horoshop не
зможе забрати їх за прямими URL. Скрипт качає їх у photos/ (<артикул>.<ext>,
<артикул>@N.<ext>) і пише в "Фото"/"Галерея" імена файлів. Якщо задано
PHOTO_BASE_URL — пише повні URL (https://<домен>/<шлях>/<файл>) під
завантажені туди фото.

Запуск:  python build_chuyka_import.py [PHOTO_BASE_URL]
Потребує: openpyxl
"""
import openpyxl, re, html, os, sys, urllib.request
import xml.etree.ElementTree as ET
from openpyxl import Workbook

FEED  = "marketplace_feed_29_09 (1).xml"
PRICE = "ПРАЙС_ДЛЯ_ПАРТНЕРІВ_ОНОВЛЕНИЙ_0110.xlsx"
OUT   = "horoshop_chuyka_import.xlsx"

PHOTO_DIR = "photos"
PHOTO_BASE_URL = (sys.argv[1].rstrip("/") + "/") if len(sys.argv) > 1 else ""
UA_NAME_FIX = {  # у фіді "Чуйка 3 0" (втрачена крапка); назву беремо як у прайсі
    "BB-19441": "Кріплення для детектора дронів «Чуйка»",
}
SECTION = "Детектори дронів"   # назва розділу на сайті (колонка "Раздел")
BRAND   = "Blue Bird"
CURRENCY = "UAH"

# 7 позицій «Чуйка» з прайсу → ключ для пошуку у фіді (підрядок у name_ua)
TARGETS = [
    ("Автомобільний тримач для «Чуйка»",                      "Автомобільний тримач"),
    ("Виносна антена для дрон детектора «Чуйка» (10 м)",       "Чуйка» 10 м"),
    ("Виносна антена для дрон детектора «Чуйка» (3 м кабелі)", "Чуйка» 3 м"),
    ("Виносна антена для дрон детектора «Чуйка» (5 м кабелі)", "Чуйка» 5 м"),
    ("Кріплення для детектора дронів «Чуйка»",                 "Кріплення для фіксації"),
    ("Пристрій радіоелектронної розвідки «Чуйка 3.0»",         "Чуйка 3.0»"),
    ("Пристрій радіоелектронної розвідки «Чуйка 4.0»",         "Чуйка 4.0»"),
]

HEADER = [
    "Артикул","Родительский артикул","Артикул для отображения на сайте",
    "Название модификации (UA)","Название модификации (RU)",
    "Название (UA)","Название (RU)","Бренд","Раздел","Цена","Старая цена",
    "Валюта","Отображать","Наличие","Дополнительные разделы","Фото","Галерея",
    "Обзор 360","Алиас","Ссылка","Дата добавления","HTML title (UA)",
    "HTML title (RU)","META keywords (UA)","META keywords (RU)",
    "META description (UA)","META description (RU)","h1 заголовок (UA)",
    "h1 заголовок (RU)","Поставщик","Иконки","Скидка %","Популярность",
    "Количество","Описание товара (UA)","Описание товара (RU)",
    "Короткое описание (UA)","Короткое описание (RU)","Цвет","Тип гарантии",
    "Гарантийный срок, мес.","Дата и время окончания акции","Текст акции (UA)",
    "Текст акции (RU)","Описание для маркетплейсов (UA)",
    "Описание для маркетплейсов (RU)","Выгружать на маркетплейсы",
    "Состояние товара","Только для взрослых","Срок доставки предзаказа в днях",
    "Производитель и модель","Материал","Материал (на английском)","Вес",
    "Тип упаковки","Код страны производителя","Единицы измерения",
    "Код УКТ ВЭД","«Оплата частями» ПриватБанка","«Покупка частями» от monobank",
    "Код налоговой группы Вчасно.Каса","Уникальный код налога","Штрихкод",
    "Код производителя товара (MPN)","На складе для Prom","Электронный товар",
]


def fetch(url, dest):
    req = urllib.request.Request(url, headers={
        "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
                      "(KHTML, like Gecko) Chrome/120 Safari/537.36",
        "Referer": "https://www.blue-bird.tech/",
        "Accept": "image/jpeg,image/png,*/*"})
    with urllib.request.urlopen(req, timeout=60) as r:
        data = r.read()
    with open(dest, "wb") as f:
        f.write(data)


def photo_names(sku, urls):
    """Качає фото в PHOTO_DIR, повертає імена файлів у порядку urls."""
    os.makedirs(PHOTO_DIR, exist_ok=True)
    names = []
    for i, u in enumerate(urls):
        ext = os.path.splitext(u.split("?")[0])[1].lower() or ".jpg"
        ext = ".jpg" if ext == ".jpeg" else ext
        name = f"{sku}{'' if i == 0 else '@%d' % i}{ext}"
        dest = os.path.join(PHOTO_DIR, name)
        if not os.path.exists(dest):
            fetch(u, dest)
        names.append(PHOTO_BASE_URL + name)
    return names


def strip_hotlinks(text):
    """Прибирає <img> з blue-bird.tech в описі (хотлінк віддає 403)."""
    return re.sub(r'<img[^>]*blue-bird\.tech[^>]*>', '', text)


def load_feed():
    root = ET.parse(FEED).getroot()
    offers = root.find("shop").find("offers")
    out = []
    for o in offers:
        out.append(o)
    return out


def find_offer(offers, key):
    for o in offers:
        name = (o.findtext("name_ua") or o.findtext("name") or "")
        if key in name:
            return o
    return None


def slugify(text):
    tr = {
        'а':'a','б':'b','в':'v','г':'h','ґ':'g','д':'d','е':'e','є':'ie','ж':'zh',
        'з':'z','и':'y','і':'i','ї':'i','й':'i','к':'k','л':'l','м':'m','н':'n',
        'о':'o','п':'p','р':'r','с':'s','т':'t','у':'u','ф':'f','х':'kh','ц':'ts',
        'ч':'ch','ш':'sh','щ':'shch','ь':'','ю':'iu','я':'ia',
    }
    s = text.lower()
    s = re.sub(r'[«»"\',.()]', '', s)
    out = ''.join(tr.get(ch, ch) for ch in s)
    out = re.sub(r'[^a-z0-9]+', '-', out).strip('-')
    return out[:80]


def params_to_html(offer):
    rows = []
    for p in offer.findall("param"):
        k = p.get("name", "").strip()
        v = (p.text or "").strip()
        if k and v:
            rows.append(f"<tr><td>{html.escape(k)}</td><td>{html.escape(v)}</td></tr>")
    if not rows:
        return ""
    return ("<p><strong>Характеристики:</strong></p>"
            "<table>" + "".join(rows) + "</table>")


def build_desc(offer, lang):
    base = offer.findtext("description_ua") if lang == "ua" else offer.findtext("description")
    base = (base or "").strip()
    out = base
    if lang == "ua":
        out = (base + params_to_html(offer)) if base else params_to_html(offer)
    return out


def main():
    offers = load_feed()

    # прайс -> {назва: ціна}
    wb_p = openpyxl.load_workbook(PRICE, read_only=True, data_only=True)
    ws = wb_p.active
    price_rows = list(ws.iter_rows(values_only=True))
    price_hdr = price_rows[0]
    pi = {h: i for i, h in enumerate(price_hdr)}
    price_map = {}
    for r in price_rows[1:]:
        if not r or r[pi["Номенклатура"]] is None:
            continue
        nom = str(r[pi["Номенклатура"]]).strip()
        price_map[nom] = r[pi["ЦІНА З ПДВ"]]

    wb = Workbook(); out = wb.active; out.title = "Export"
    out.append(HEADER)

    count = 0
    missing = []
    for price_name, feed_key in TARGETS:
        offer = find_offer(offers, feed_key)
        price = price_map.get(price_name)
        if offer is None:
            missing.append(("no feed", price_name)); continue
        if price is None:
            missing.append(("no price", price_name)); continue

        vcode   = offer.findtext("vendorCode") or f"BB-{offer.get('id')}"
        name_ua = (UA_NAME_FIX.get(vcode) or offer.findtext("name_ua")
                   or offer.findtext("name") or price_name)
        name_ru = offer.findtext("name") or name_ua
        country = offer.findtext("country_of_origin") or "Україна"
        pics    = [p.text for p in offer.findall("picture") if p.text]
        files    = photo_names(vcode, pics)
        main_pic = files[0] if files else ""
        gallery  = ";".join(files[1:])
        avail = "В наявності" if offer.get("available") == "true" else "Нет в наличии"

        row = {h: "" for h in HEADER}
        row["Артикул"] = vcode
        row["Название (UA)"] = name_ua
        row["Название (RU)"] = name_ru
        row["Бренд"] = BRAND
        row["Раздел"] = SECTION
        row["Цена"] = price
        row["Валюта"] = CURRENCY
        row["Отображать"] = "Да"
        row["Наличие"] = avail
        row["Фото"] = main_pic
        row["Галерея"] = gallery
        row["Алиас"] = slugify(name_ua)
        row["HTML title (UA)"] = f"{name_ua} — купити, ціна в Україні"
        row["META description (UA)"] = (
            f"{name_ua} від Blue Bird Tech. Детектор дронів «Чуйка». "
            f"Офіційний партнер, гарантія, доставка по Україні."
        )[:255]
        row["h1 заголовок (UA)"] = name_ua
        row["HTML title (RU)"] = f"{name_ru} — купить, цена в Украине"
        row["META description (RU)"] = (
            f"{name_ru} от Blue Bird Tech. Детектор дронов «Чуйка». "
            f"Официальный партнёр, гарантия, доставка по Украине."
        )[:255]
        row["h1 заголовок (RU)"] = name_ru
        row["Короткое описание (RU)"] = name_ru
        row["Описание товара (UA)"] = strip_hotlinks(build_desc(offer, "ua"))
        row["Описание товара (RU)"] = strip_hotlinks(build_desc(offer, "ru"))
        row["Короткое описание (UA)"] = name_ua
        row["Количество"] = 1 if avail == "В наявності" else 0
        row["Состояние товара"] = "Новий"
        row["Только для взрослых"] = "Нет"
        row["Код страны производителя"] = country
        row["Производитель и модель"] = name_ua
        row["«Оплата частями» ПриватБанка"] = "Выкл"
        row["«Покупка частями» от monobank"] = "Выкл"
        row["На складе для Prom"] = "Да" if avail == "В наявності" else "Нет"
        row["Электронный товар"] = "Нет"
        row["Код производителя товара (MPN)"] = vcode

        out.append([row[h] for h in HEADER])
        count += 1

    wb.save(OUT)
    print(f"OK: {count}/7 товарів -> {OUT}")
    if missing:
        print("УВАГА, не зматчено:")
        for why, n in missing:
            print(f"  [{why}] {n}")

if __name__ == "__main__":
    main()
