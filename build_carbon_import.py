# -*- coding: utf-8 -*-
"""
Готує Horoshop-імпорт (66 колонок) з carbon_net_ua_catalog.xlsx.
Один універсальний файл: колонка "Раздел" = назва нової категорії,
яку ви створюєте в адмінці за інструкцією. Без варіацій — кожен товар
окрема картка, щоб файл заливався на кілька сайтів "як є".

Запуск:  python build_carbon_import.py
Потребує: openpyxl
"""
import openpyxl, re, html
from openpyxl import Workbook

SRC = "carbon_net_ua_catalog.xlsx"
OUT = "horoshop_carbon_import.xlsx"

# --- НАЛАШТУВАННЯ (міняйте під свої сайти) ---------------------------------
SECTION = "Карбон"          # назва розділу/категорії на сайті (колонка "Раздел")
BRAND   = "Carbon"          # бренд; лишіть порожнім "" якщо не треба
CURRENCY = "UAH"
# ---------------------------------------------------------------------------

# 66 колонок офіційного Horoshop-експорту (порядок важливий)
HEADER = [
    "Артикул","Родительский артикул","Артикул для отображения на сайте",
    "Название модификации (UA)","Название модификации (RU)",
    "Название (UA)","Название (RU)","Бренд","Раздел","Цена","Старая цена",
    "Валюта","Отображать","Наличие","Дополнительные разделы","Фото","Галерея",
    "Обзор 360","Алиас","Ссылка","Дата добавления","HTML title (UA)",
    "HTML title (RU)","META keywords (UA)","META keywords (RU)",
    "META description (UA)","META description (RU)","h1 заголовок (UA)",
    "h1 заголовок (RU)","Поставщик","Иконки","Скидка %","Популярность",
    "Количество","Описание товара (UA)","Описание товара (RU)",
    "Короткое описание (UA)","Короткое описание (RU)","Цвет","Тип гарантии",
    "Гарантийный срок, мес.","Дата и время окончания акции","Текст акции (UA)",
    "Текст акции (RU)","Описание для маркетплейсов (UA)",
    "Описание для маркетплейсов (RU)","Выгружать на маркетплейсы",
    "Состояние товара","Только для взрослых","Срок доставки предзаказа в днях",
    "Производитель и модель","Материал","Материал (на английском)","Вес",
    "Тип упаковки","Код страны производителя","Единицы измерения",
    "Код УКТ ВЭД","«Оплата частями» ПриватБанка","«Покупка частями» от monobank",
    "Код налоговой группы Вчасно.Каса","Уникальный код налога","Штрихкод",
    "Код производителя товара (MPN)","На складе для Prom","Электронный товар",
]

AVAIL_MAP = {
    "В наявності": "В наявності",
    "Очікується": "Нет в наличии",
    "Очікується (СКОРО В НАЯВНОСТІ)": "Нет в наличии",
}

def parse_price(s):
    """'13500 грн (стара ціна 15600 грн, знижка -13%)' -> (13500, 15600, 13)"""
    s = str(s or "")
    nums = re.findall(r"(\d[\d\s]*)\s*грн", s)
    price = old = 0
    if nums:
        price = int(nums[0].replace(" ", ""))
        if len(nums) > 1:
            old = int(nums[1].replace(" ", ""))
    disc = re.search(r"-(\d+)\s*%", s)
    discount = int(disc.group(1)) if disc else 0
    return price, old, discount

def slugify(url, art):
    m = re.search(r"carbon\.net\.ua/([^/]+)/?", str(url or ""))
    if m:
        return m.group(1)
    return f"carbon-{art}"

def specs_to_html(specs):
    """'Розмір 1000x1000 мм; Вага 3400г; ...' -> <ul><li>..</li></ul>"""
    parts = [p.strip() for p in str(specs or "").split(";") if p.strip()]
    if not parts:
        return ""
    lis = "".join(f"<li>{html.escape(p)}</li>" for p in parts)
    return f"<ul>{lis}</ul>"

def get_spec(specs, key):
    m = re.search(rf"{key}\s+([^;]+)", str(specs or ""))
    return m.group(1).strip() if m else ""

def build_description(name, desc, specs):
    out = []
    if desc and str(desc) != "None":
        out.append(f"<p>{html.escape(str(desc).strip())}</p>")
    sp = specs_to_html(specs)
    if sp:
        out.append("<p><b>Характеристики:</b></p>")
        out.append(sp)
    if not out:
        out.append(f"<p>{html.escape(str(name))}</p>")
    return "\n".join(out)

def main():
    wb_in = openpyxl.load_workbook(SRC, read_only=True)
    ws = wb_in.active
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]; data = rows[1:]
    idx = {h: i for i, h in enumerate(hdr)}

    wb = Workbook(); out = wb.active; out.title = "Export"
    out.append(HEADER)

    seen = set()
    count = 0
    for r in data:
        g = lambda k: r[idx[k]] if k in idx else None
        name = str(g("Назва товару") or "").strip()
        art  = str(g("Артикул") or "").strip()
        if not name or not art:
            continue
        if art in seen:          # на випадок дублів артикулів
            art = f"{art}-{count}"
        seen.add(art)

        price, old, disc = parse_price(g("Ціна"))
        avail = AVAIL_MAP.get(str(g("Наявність") or "").strip(), "Нет в наличии")
        specs = g("Характеристики")
        color = get_spec(specs, "Колір")
        material = "Карбон (вуглепластик)"
        alias = slugify(g("URL сторінки"), art)
        desc_html = build_description(name, g("Опис"), specs)

        row = {h: "" for h in HEADER}
        row["Артикул"] = art
        row["Название (UA)"] = name
        row["Название (RU)"] = name
        row["Бренд"] = BRAND
        row["Раздел"] = SECTION
        row["Цена"] = price
        row["Старая цена"] = old if old else ""
        row["Валюта"] = CURRENCY
        row["Отображать"] = "Да"
        row["Наличие"] = avail
        row["Алиас"] = alias
        row["Скидка %"] = disc if disc else ""
        row["Количество"] = 1 if avail == "В наявності" else 0
        row["HTML title (UA)"] = f"{name} — купити в Україні"
        row["META description (UA)"] = (
            f"{name}. Карбонові вироби для FPV-дронів та RC-моделей. "
            f"Швидка доставка по Україні."
        )[:255]
        row["h1 заголовок (UA)"] = name
        row["Описание товара (UA)"] = desc_html
        row["Описание товара (RU)"] = desc_html
        row["Короткое описание (UA)"] = name
        row["Цвет"] = color
        row["Состояние товара"] = "Новий"
        row["Только для взрослых"] = "Нет"
        row["Материал"] = material
        row["«Оплата частями» ПриватБанка"] = "Выкл"
        row["«Покупка частями» от monobank"] = "Выкл"
        row["На складе для Prom"] = "Да" if avail == "В наявності" else "Нет"
        row["Электронный товар"] = "Нет"

        out.append([row[h] for h in HEADER])
        count += 1

    wb.save(OUT)
    print(f"OK: {count} товарів -> {OUT}")
    print(f"Розділ у файлі: '{SECTION}'  |  Бренд: '{BRAND}'")

if __name__ == "__main__":
    main()
