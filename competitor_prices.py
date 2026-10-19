#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
competitor_prices.py — моніторинг цін конкурентів БЕЗ браузера.

Парсить ціни товарів з трьох сайтів-конкурентів через curl_cffi
(імітація TLS-відбитку Chrome) + власний розв'язувач anti-bot challenge.
Браузер (claude-in-chrome / Selenium / Playwright) не потрібен.

Залежності:  pip install curl_cffi beautifulsoup4

Кожен сайт має свій захист — реалізовано окремо:
  • dron-shop.com.ua  — adm.tools proof-of-work (hashcash SHA-256).
                        Відтворено повний flow fetch→challenge→solve→confirm.
  • ekran.store       — Horoshop JS proof-of-work, що лише ставить статичну
                        cookie challenge_passed=<hash>. Підставляємо її напряму.
  • modelistam.com.ua — без захисту, звичайний GET.

Запуск:   python3 competitor_prices.py "Matek"
"""
import sys, re, json, time, struct, hashlib
from curl_cffi import requests as cf
from bs4 import BeautifulSoup

UA_LANG = {"Accept-Language": "uk-UA,uk;q=0.9,ru;q=0.8,en;q=0.7"}

# ---------------------------------------------------------------------------
# dron-shop.com.ua — adm.tools proof-of-work
# ---------------------------------------------------------------------------
class AdmToolsSession:
    """curl_cffi-сесія, що проходить adm.tools PoW-challenge і далі ходить як браузер."""

    def __init__(self, base, impersonate="chrome124"):
        self.base = base.rstrip("/") + "/"
        self.s = cf.Session(impersonate=impersonate)
        self._metrics = {"mouse": 5, "keyboard": 2, "focus": 1, "platform": "Win32",
                         "memory": 8, "cpu": 8, "connection": {}, "plugins": ["PDF Viewer"]}
        self._solved = False

    def _rpc(self, action, **kw):
        data = {"action": action, "metrics": self._metrics, **kw}
        return self.s.post(self.base, json=data, timeout=40, headers={
            "Content-Type": "application/json",
            "X-Requested-With": "XMLHttpRequest",
            "Referer": self.base,
            "Origin": self.base.rstrip("/"),
            **UA_LANG,
        })

    class CaptchaRequired(RuntimeError):
        """adm.tools ескалював до візуальної CAPTCHA (SVG) — потрібна людина."""
        def __init__(self, svg):
            super().__init__("adm.tools повернув візуальну CAPTCHA замість PoW")
            self.svg = svg

    @staticmethod
    def _bits_zero(digest, bits):
        full, rem = divmod(bits, 8)
        if any(digest[i] for i in range(full)):
            return False
        if rem and (digest[full] & ((0xff << (8 - rem)) & 0xff)):
            return False
        return True

    def _solve_pow(self, nonce, bits, solutions):
        prefix = (nonce + ":").encode()
        counters, counter = [], 0
        while len(counters) < solutions:
            if self._bits_zero(hashlib.sha256(prefix + struct.pack(">Q", counter)).digest(), bits):
                counters.append(counter)
            counter += 1
        return counters

    def solve(self):
        """Пройти challenge один раз (ідемпотентно)."""
        if self._solved:
            return
        self.s.get(self.base, timeout=30, headers=UA_LANG)   # 429 + warm-up cookie
        self._rpc("fetch", auto=True)                        # віддає sl_sess
        ch = self._rpc("challenge").json()
        # якщо сервер уже вважає нас пройденими — challenge віддає {"action":"switch"}
        if ch.get("action") == "switch":
            self._solved = True
            return
        # при частих запитах adm.tools ескалює: замість PoW віддає SVG-CAPTCHA,
        # яку чистим Python не розв'язати — пробрасываем нагору, щоб людина пройшла в браузері
        if "challenge" in ch and "nonce" not in ch and "<svg" in str(ch.get("challenge", "")):
            raise AdmToolsSession.CaptchaRequired(ch["challenge"])
        if "nonce" not in ch:
            self._solved = True
            return
        t0 = time.time()
        counters = self._solve_pow(ch["nonce"], ch["bits"], ch["solutions"])
        self._rpc("confirm", id=ch["id"], counters=counters, time=int(time.time() - t0))
        self._solved = True

    def get(self, url):
        self.solve()
        return self.s.get(url, timeout=40, headers={"Referer": self.base,
                          "X-Requested-With": "XMLHttpRequest", **UA_LANG})


def scrape_dronshop(query, retries=3):
    url = f"https://dron-shop.com.ua/index.php?route=product/search&search={query}&limit=100"
    soup = None
    for attempt in range(retries):
        sess = AdmToolsSession("https://dron-shop.com.ua/")   # свіжа сесія = свіжий challenge
        html = sess.get(url).text
        soup = BeautifulSoup(html, "html.parser")
        if soup.select(".product-thumb"):
            break
        # adm.tools ескалує до візуальної CAPTCHA при частих запитах —
        # пауза між спробами, щоб не тригерити ускладнення
        time.sleep(8.0)
    out = []
    for item in soup.select(".product-thumb"):
        name_el = item.select_one(".h2 a, .caption h4 a, .name a, h4 a")
        price_el = item.select_one(".price-new, .price")
        if not name_el or not price_el:
            continue
        name = name_el.get_text(strip=True)
        # розмітка .price:  "2995 <span>3695</span>"  — поточна ціна текстом,
        # закреслена стара у вкладеному <span>.
        old_el = price_el.find("span")
        old = int(re.sub(r"\D", "", old_el.get_text())) if old_el and old_el.get_text(strip=True) else None
        # поточна ціна = текст .price без вмісту <span>
        if old_el:
            old_el.extract()
        cur_txt = re.sub(r"\D", "", price_el.get_text())
        if not cur_txt:
            continue
        out.append({"name": name, "price": int(cur_txt), "old_price": old})
    return out


# ---------------------------------------------------------------------------
# ekran.store — Horoshop + статична challenge-cookie
# ---------------------------------------------------------------------------
# Хеш challenge_passed той самий для всього домену (міняється зрідка —
# якщо почне віддавати 404/порожнє, зняти свіжий із <script> головної).
EKRAN_COOKIE = "14c5cabb3604a4dbfc6352e88e052f12ebbd9b44cab27b3e60bb4d12c3a68541"


def ekran_session():
    s = cf.Session(impersonate="chrome124")
    s.cookies.set("challenge_passed", EKRAN_COOKIE, domain="ekran.store")
    return s


def scrape_ekran_product(url, sess=None):
    """Ціна + наявність з картки товару Horoshop (JSON-LD / itemprop)."""
    sess = sess or ekran_session()
    html = sess.get(url, timeout=30, headers=UA_LANG).text
    price = re.search(r'itemprop="price"[^>]*content="([\d.]+)"', html) \
        or re.search(r'"price"\s*:\s*"?([\d.]+)', html)
    avail = re.search(r'(InStock|OutOfStock|PreOrder)', html)
    return {"url": url,
            "price": float(price.group(1)) if price else None,
            "availability": avail.group(1) if avail else None}


# ---------------------------------------------------------------------------
# modelistam.com.ua — без захисту
# ---------------------------------------------------------------------------
def scrape_modelistam(query):
    s = cf.Session(impersonate="chrome124")
    url = f"https://modelistam.com.ua/ua/advanced_search_result.php?keywords={query}"
    html = s.get(url, timeout=40, headers=UA_LANG).text
    out = []
    for card in re.split(r'<span class="product-thumb__title-inner"', html)[1:]:
        tm = re.search(r'title="([^"]*)"', card)
        cm = re.search(r'product-thumb__code-value">([^<]*)<', card)
        pm = re.search(r'product-thumb__price[^"]*">(.*?)</', card, re.S)
        price_txt = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", pm.group(1))).strip() if pm else ""
        digits = re.sub(r"\D", "", price_txt)
        out.append({"code": cm.group(1).strip() if cm else "",
                    "name": tm.group(1) if tm else "",
                    "price": int(digits) if digits else None})
    return out


# ---------------------------------------------------------------------------
# ekran.store — пошук за запитом через catalog-sitemap
# ---------------------------------------------------------------------------
def scrape_ekran(query):
    """Знайти товари бренда/запиту в ekran.store і зняти ціну+наявність з карток."""
    sess = ekran_session()
    sm = sess.get("https://ekran.store/content/export/ekran.store/catalog-sitemap.xml",
                  timeout=60, headers=UA_LANG).text
    urls = [u for u in re.findall(r"<loc>([^<]*)</loc>", sm)
            if query.lower() in u.lower() and "/ru/" not in u]
    out = []
    for u in urls:
        d = scrape_ekran_product(u, sess)
        d["url"] = u
        out.append(d)
    return out


# ---------------------------------------------------------------------------
# Експорт у Excel
# ---------------------------------------------------------------------------
def export_excel(path, our_rows, header_cols, data_rows, sheet="Порівняння"):
    """our_rows/data_rows — список словників; header_cols — порядок колонок."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    hdr_fill = PatternFill("solid", fgColor="0E7490")
    hdr_font = Font(bold=True, color="FFFFFF")
    for c, name in enumerate(header_cols, 1):
        cell = ws.cell(1, c, name)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(data_rows, 2):
        for c, key in enumerate(header_cols, 1):
            ws.cell(r, c, row.get(key))
    for col in ws.columns:
        w = max((len(str(x.value)) for x in col if x.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 10), 60)
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


if __name__ == "__main__":
    q = sys.argv[1] if len(sys.argv) > 1 else "Matek"
    print(f"# Пошук '{q}' по конкурентах (без браузера)\n")

    print("== dron-shop.com.ua (adm.tools PoW) ==")
    try:
        for p in scrape_dronshop(q):
            print(f"  {p['price']:>7} | {p['name']}")
    except AdmToolsSession.CaptchaRequired as e:
        open("dronshop_captcha.svg", "w", encoding="utf-8").write(e.svg)
        print("  !! adm.tools віддав CAPTCHA — збережено в dronshop_captcha.svg")
        print("     Відкрий dron-shop.com.ua у браузері вручну та пройди перевірку.")

    print("\n== modelistam.com.ua ==")
    for p in scrape_modelistam(q):
        if q.lower() in (p["code"] + p["name"]).lower():
            print(f"  {str(p['price']):>7} | {p['code']:20} | {p['name']}")

    print("\n== ekran.store ==")
    for p in scrape_ekran(q):
        print(f"  {str(p['price']):>7} | {p['availability']:10} | {p['url']}")
